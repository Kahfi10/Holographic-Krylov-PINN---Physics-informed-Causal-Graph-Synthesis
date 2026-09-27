"""
latex_to_word_omml.py
Helper utility to convert LaTeX expressions into Microsoft Word Native Office Math (<m:oMath>) elements.
Uses latex2mathml and Microsoft's official MML2OMML.XSL stylesheet.
"""

import os
import re
from latex2mathml.converter import convert as latex2mathml_convert
from lxml import etree
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Default XSLT path on Windows Microsoft Office installation
DEFAULT_XSLT_PATH = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"

class LatexToOmmlConverter:
    def __init__(self, xslt_path=DEFAULT_XSLT_PATH):
        if not os.path.exists(xslt_path):
            raise FileNotFoundError(f"MML2OMML.XSL not found at: {xslt_path}")
        xslt_doc = etree.parse(xslt_path)
        self.transform = etree.XSLT(xslt_doc)

    def to_omml(self, latex_code: str):
        """Convert a single LaTeX string into an lxml OMML root element (<m:oMath>)."""
        clean_latex = latex_code.strip()
        mathml = latex2mathml_convert(clean_latex)
        mathml_tree = etree.fromstring(mathml)
        omml_tree = self.transform(mathml_tree)
        return omml_tree.getroot()

    def add_inline_math_runs(self, paragraph, text: str, font_name="Cambria", size_pt=11, bold=False, italic=False):
        """
        Parses text containing inline LaTeX math wrapped in `$ ... $` and appends
        native Office Math (<m:oMath>) elements for math and standard formatted runs for plain text.
        """
        parts = re.split(r'(\$[^\$]+\$)', text)
        for part in parts:
            if part.startswith('$') and part.endswith('$') and len(part) > 2:
                latex_code = part[1:-1]
                try:
                    omml = self.to_omml(latex_code)
                    paragraph._p.append(omml)
                except Exception as e:
                    r = paragraph.add_run(latex_code)
                    r.font.name = font_name
                    r.font.size = Pt(size_pt)
                    r.italic = True
            else:
                if part:
                    r = paragraph.add_run(part)
                    r.font.name = font_name
                    r.font.size = Pt(size_pt)
                    r.bold = bold
                    r.italic = italic

    def add_display_equation(self, doc, latex_code: str, eq_num: str, center_dxa=4819, right_dxa=9638, font_name="Cambria", size_pt=11):
        """
        Adds a numbered display equation paragraph to the document with:
        - Center-aligned equation via tab stop at center_dxa
        - Right-aligned equation number `(eq_num)` via tab stop at right_dxa
        """
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0

        pPr = p._p.get_or_add_pPr()
        tabsXml = parse_xml(
            f'<w:tabs {nsdecls("w")}>'
            f'<w:tab w:val="center" w:pos="{center_dxa}"/>'
            f'<w:tab w:val="right" w:pos="{right_dxa}"/>'
            f'</w:tabs>'
        )
        pPr.append(tabsXml)

        p.add_run("\t")
        try:
            omml = self.to_omml(latex_code)
            p._p.append(omml)
        except Exception:
            r_err = p.add_run(latex_code)
            r_err.font.name = font_name
            r_err.font.size = Pt(size_pt)
            r_err.italic = True

        r_num = p.add_run(f"\t({eq_num})")
        r_num.font.name = font_name
        r_num.font.size = Pt(size_pt)
        r_num.bold = False
        return p
